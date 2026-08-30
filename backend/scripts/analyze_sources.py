import openpyxl
from collections import defaultdict

CRM_PATH = 'c:/Users/user/Alpha-Pro-Mena-CRM\\CRM.xlsx'
GULF_PATH = 'c:/Users/user/Alpha-Pro-Mena-CRM\\Gulf Leads .xlsx'

# ========================
# CRM.xlsx - NO HEADERS, data from row 1
# Col 0=Name, 1=Company, 2=Position, 3=Phone, 4=Email, 5=SalesPerson
# ========================
print("=" * 60)
print("CRM.xlsx FULL ANALYSIS")
print("=" * 60)

wb_crm = openpyxl.load_workbook(CRM_PATH, read_only=True, data_only=True)
ws = wb_crm['Sheet1']

crm_by_rep = defaultdict(list)
crm_total = 0
crm_blank_rows = 0

for row_idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
    # Skip completely blank rows
    if all(v is None or str(v).strip() == '' for v in row):
        crm_blank_rows += 1
        continue
    
    name = row[0] if row[0] is not None else None
    company = row[1] if len(row) > 1 and row[1] is not None else None
    position = row[2] if len(row) > 2 and row[2] is not None else None
    phone = str(row[3]).strip() if len(row) > 3 and row[3] is not None else None
    email = row[4] if len(row) > 4 and row[4] is not None else None
    sp = str(row[5]).strip() if len(row) > 5 and row[5] is not None else 'UNKNOWN'
    
    crm_by_rep[sp].append({
        'row_idx': row_idx,
        'name': name,
        'company': company,
        'position': position,
        'phone': phone,
        'email': email,
    })
    crm_total += 1

wb_crm.close()

print(f"\nBlank rows skipped: {crm_blank_rows}")
print(f"Total data rows: {crm_total}")
print(f"\nBy Sales Person:")
for rep in sorted(crm_by_rep.keys()):
    rows = crm_by_rep[rep]
    print(f"  '{rep}': {len(rows)}")
    # Show first 2
    for r in rows[:2]:
        print(f"    row={r['row_idx']} | {r['name']} | {r['company']} | {r['phone']}")

# ========================
# Gulf Leads - Leads sheet
# Row 1 = headers (col 0=Name, 1=Company, 2=Position, 3=Phone, 4=Email, 5=SalesPerson)
# ========================
print()
print("=" * 60)
print("Gulf Leads .xlsx - LEADS SHEET FULL ANALYSIS")
print("=" * 60)

wb_gulf = openpyxl.load_workbook(GULF_PATH, read_only=True, data_only=True)
ws_leads = wb_gulf['Leads']

gulf_by_rep = defaultdict(list)
gulf_total = 0
gulf_blank_rows = 0
skip_header = True

for row_idx, row in enumerate(ws_leads.iter_rows(values_only=True), start=1):
    if skip_header:
        skip_header = False
        print(f"Header row: {row}")
        continue
    
    # Skip completely blank rows
    if all(v is None or str(v).strip() == '' for v in row):
        gulf_blank_rows += 1
        continue
    
    name = row[0] if row[0] is not None else None
    company = row[1] if len(row) > 1 and row[1] is not None else None
    position = row[2] if len(row) > 2 and row[2] is not None else None
    phone = str(row[3]).strip() if len(row) > 3 and row[3] is not None else None
    email = row[4] if len(row) > 4 and row[4] is not None else None
    sp = str(row[5]).strip() if len(row) > 5 and row[5] is not None else 'UNKNOWN'
    attempts_1 = row[6] if len(row) > 6 and row[6] is not None else None
    attempts_2 = row[7] if len(row) > 7 and row[7] is not None else None
    attempts_3 = row[8] if len(row) > 8 and row[8] is not None else None
    notes = row[10] if len(row) > 10 and row[10] is not None else None
    
    gulf_by_rep[sp].append({
        'row_idx': row_idx,
        'name': name,
        'company': company,
        'position': position,
        'phone': phone,
        'email': email,
        'attempt_1': attempts_1,
        'attempt_2': attempts_2,
        'attempt_3': attempts_3,
        'notes': notes,
    })
    gulf_total += 1

wb_gulf.close()

print(f"\nBlank rows skipped: {gulf_blank_rows}")
print(f"Total data rows (Leads sheet only): {gulf_total}")
print(f"\nBy Sales Person:")
for rep in sorted(gulf_by_rep.keys()):
    rows = gulf_by_rep[rep]
    print(f"  '{rep}': {len(rows)}")
    # Show first 2
    for r in rows[:2]:
        print(f"    row={r['row_idx']} | {r['name']} | {r['company']} | {r['phone']}")

# Key reps we care about
print()
print("=== KEY REPS SUMMARY ===")
for crm_rep, archive_rep in [('Saleh', 'Saleh'), ('Amin', 'Amin'), ('Hassan', 'Hasan'), ('Ghaida', 'Ghaida')]:
    crm_count = len(crm_by_rep.get(crm_rep, []))
    gulf_count = len(gulf_by_rep.get(archive_rep, []))
    print(f"  CRM '{crm_rep}': {crm_count} | Gulf '{archive_rep}': {gulf_count}")
