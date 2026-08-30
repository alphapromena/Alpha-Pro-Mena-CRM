import openpyxl

wb = openpyxl.load_workbook('c:/Users/user/Alpha-Pro-Mena-CRM/Gulf Leads .xlsx', read_only=True)
for name in wb.sheetnames:
    ws = wb[name]
    rows = list(ws.iter_rows(max_row=3, values_only=True))
    print(f"=== Sheet: [{name}] ===")
    for r in rows:
        print(" ", [str(c)[:30].encode('ascii', 'replace').decode('ascii') if c is not None else None for c in r[:8]])
