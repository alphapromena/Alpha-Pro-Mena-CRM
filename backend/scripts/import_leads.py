"""
import_leads.py — Idempotent Gulf Leads importer for Alpha Pro MENA CRM.

Usage (from backend/ directory):
    python -m scripts.import_leads                    # full import
    python -m scripts.import_leads --dry-run          # preview counts, no writes
    python -m scripts.import_leads --sheet Leads      # single sheet
    python -m scripts.import_leads --file /path/to/file.xlsx

Environment:
    LEADS_XLSX_PATH=<path>  — override default workbook search path

The workbook is searched in this order:
  1. --file argument
  2. LEADS_XLSX_PATH env var
  3. <repo_root>/Gulf Leads  (1).xlsx   (two spaces — actual filename)
  4. <repo_root>/Gulf Leads (1).xlsx    (one space variant)
  5. <repo_root>/data/gulf_leads_source.xlsx  (legacy location)
"""
import argparse
import asyncio
import os
import sys
from pathlib import Path
from typing import Dict, List, Optional

_BACKEND_DIR = Path(__file__).resolve().parent.parent
if str(_BACKEND_DIR) not in sys.path:
    sys.path.insert(0, str(_BACKEND_DIR))

# Ensure Windows stdout/stderr handles Arabic and Unicode characters cleanly
if sys.stdout and hasattr(sys.stdout, "reconfigure"):
    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass
if sys.stderr and hasattr(sys.stderr, "reconfigure"):
    try:
        sys.stderr.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass

import logging
logging.getLogger("sqlalchemy.engine").setLevel(logging.WARNING)

_REPO_ROOT = _BACKEND_DIR.parent
_DEFAULT_XLSX_CANDIDATES = [
    _REPO_ROOT / "Gulf Leads  (1).xlsx",
    _REPO_ROOT / "Gulf Leads (1).xlsx",
    _REPO_ROOT / "data" / "gulf_leads_source.xlsx",
]


def _find_xlsx(args_file: Optional[str]) -> Optional[Path]:
    if args_file:
        p = Path(args_file)
        return p if p.exists() else None
    env_path = os.environ.get("LEADS_XLSX_PATH")
    if env_path:
        p = Path(env_path)
        return p if p.exists() else None
    for candidate in _DEFAULT_XLSX_CANDIDATES:
        if candidate.exists():
            return candidate
    return None


def _parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Import Gulf Leads workbook into the Alpha Pro MENA CRM."
    )
    parser.add_argument("--dry-run", action="store_true",
                        help="Preview counts only — no DB writes.")
    parser.add_argument("--sheet", metavar="NAME", default=None,
                        help="Import a single named sheet.")
    parser.add_argument("--file", metavar="PATH", default=None,
                        help="Path to .xlsx workbook.")
    return parser.parse_args()


async def _run(args: argparse.Namespace) -> None:
    try:
        import openpyxl
    except ImportError:
        print("ERROR: openpyxl is required. Run: pip install openpyxl")
        sys.exit(1)

    from app.imports.workbook_reader import (
        GULF_LEADS_SHEET_CONFIGS,
        SKIP_SHEETS,
        parse_sheet,
        read_companies_tab,
    )
    from app.imports.db_writer import run_import

    xlsx_path = _find_xlsx(args.file)
    if xlsx_path is None or not xlsx_path.exists():
        print("\nERROR: Workbook not found.")
        for c in _DEFAULT_XLSX_CANDIDATES:
            print(f"  Searched: {c}")
        print("\nSet LEADS_XLSX_PATH or use --file <path>.\n")
        sys.exit(1)

    print(f"\nReading: {xlsx_path}")
    if args.dry_run:
        print("DRY RUN — no database writes.\n")

    wb = openpyxl.load_workbook(str(xlsx_path), read_only=True, data_only=True)
    print(f"Sheets: {wb.sheetnames}\n")

    companies_set = read_companies_tab(wb)
    print(f"Companies reference: {len(companies_set)} names loaded.\n")

    if args.sheet:
        if args.sheet not in wb.sheetnames:
            print(f"ERROR: Sheet '{args.sheet}' not in workbook.")
            sys.exit(1)
        if args.sheet in SKIP_SHEETS:
            print(f"ERROR: '{args.sheet}' is a skip sheet (not a contacts source).")
            sys.exit(1)
        sheets_to_process = [args.sheet]
    else:
        sheets_to_process = [
            s for s in wb.sheetnames
            if s not in SKIP_SHEETS and s in GULF_LEADS_SHEET_CONFIGS
        ]

    all_rows = []
    sheet_raw_counts: Dict[str, int] = {}

    for sheet_name in sheets_to_process:
        config = GULF_LEADS_SHEET_CONFIGS.get(sheet_name)
        if not config:
            print(f"  Skipping unconfigured sheet: {sheet_name!r}")
            continue
        ws = wb[sheet_name]
        rows = parse_sheet(ws, config, companies_set)
        sheet_raw_counts[sheet_name] = len(rows)
        invalid_in_sheet = sum(1 for r in rows if r.is_invalid)
        print(f"  {sheet_name!r}: {len(rows)} rows ({invalid_in_sheet} invalid)")
        all_rows.extend(rows)

    wb.close()
    print(f"\nTotal raw rows across sheets: {len(all_rows)}")

    from app.database import AsyncSessionLocal, engine, Base

    if not args.dry_run:
        async with engine.begin() as conn:
            await conn.run_sync(Base.metadata.create_all)

    async with AsyncSessionLocal() as db:
        report = await run_import(db, all_rows, dry_run=args.dry_run)
        if not args.dry_run:
            await db.commit()

    s = report.to_dict()["summary"]
    print("\n" + "=" * 65)
    print("IMPORT SUMMARY" + (" (DRY RUN)" if args.dry_run else ""))
    print("=" * 65)
    for name, cnt in sheet_raw_counts.items():
        print(f"  {name:<30}: {cnt:>5} rows")
    print(f"  {'Intra-batch duplicates merged':<30}: {s['total_intra_duplicates']:>5}")
    print(f"  {'Invalid rows':<30}: {s['total_invalid']:>5}")
    print(f"  {'Conflicts (phone/name mismatch)':<30}: {s['total_conflicts']:>5}")
    print(f"  {'Awaiting review (legacy owner)':<30}: {s['total_awaiting_review']:>5}")
    print(f"  {'Matched existing DB contacts':<30}: {s['total_db_matched']:>5}")
    print(f"  {'Genuinely new contacts':<30}: {s['total_new']:>5}")
    if not args.dry_run:
        print(f"  ---")
        print(f"  {'Inserted':<30}: {s['total_inserted']:>5}")
        print(f"  {'Updated (fields enriched)':<30}: {s['total_updated']:>5}")
        print(f"  {'Skipped (no change)':<30}: {s['total_skipped']:>5}")
        print(f"  {'Errors':<30}: {s['total_errors']:>5}")
    print("=" * 65)

    if report.invalid_rows:
        print(f"\nInvalid rows ({len(report.invalid_rows)}):")
        for r in report.invalid_rows[:10]:
            print(f"  [{r['sheet']} row {r['row']}]: {r['reason']}")
        if len(report.invalid_rows) > 10:
            print(f"  ... and {len(report.invalid_rows) - 10} more")

    if report.conflict_rows:
        print(f"\nConflicts — same phone, different name ({len(report.conflict_rows)}):")
        for r in report.conflict_rows[:10]:
            print(f"  [{r['sheet']} row {r['row']}] {r['name']}: {r['reason']}")

    if report.awaiting_review_rows:
        print(f"\nAwaiting review — unresolved salesperson ({len(report.awaiting_review_rows)}):")
        for r in report.awaiting_review_rows[:15]:
            print(f"  [{r['sheet']} row {r['row']}] {r['name']}: {r['reason']}")
        if len(report.awaiting_review_rows) > 15:
            print(f"  ... and {len(report.awaiting_review_rows) - 15} more")

    print()


def main() -> None:
    args = _parse_args()
    asyncio.run(_run(args))


if __name__ == "__main__":
    main()
