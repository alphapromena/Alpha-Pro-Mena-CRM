"""
backend/scripts/replace_active_leads_from_sheet16.py

CLI runner for authoritative CRM active data replacement:
- Active Contacts: Sourced strictly from Sheet16
- Active Companies: Authoritative master data from Companies sheet
- Leads Archive: Archived completely into Neon database leads_archive table
- Historical Demos & Follow-ups: Extracted and linked to Sheet16 contacts

Usage:
    # Dry run (safe preview without writing changes):
    python backend/scripts/replace_active_leads_from_sheet16.py --file "CRM data.xlsx" --dry-run

    # Live execution:
    python backend/scripts/replace_active_leads_from_sheet16.py --file "CRM data.xlsx"
"""
import argparse
import asyncio
import os
import sys
import openpyxl

# Add backend directory to sys.path
SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
BACKEND_DIR = os.path.dirname(SCRIPT_DIR)
if BACKEND_DIR not in sys.path:
    sys.path.insert(0, BACKEND_DIR)

from app.database import AsyncSessionLocal
from app.imports.sheet16_data_replacer import Sheet16DataReplacer


async def main():
    parser = argparse.ArgumentParser(description="Replace active CRM contacts with Sheet16 and archive Leads.")
    parser.add_argument("--file", "-f", default="CRM data.xlsx", help="Path to CRM data.xlsx workbook")
    parser.add_argument("--dry-run", action="store_true", help="Preview migration without committing changes")
    args = parser.parse_args()

    wb_path = os.path.abspath(args.file)
    if not os.path.exists(wb_path):
        print(f"ERROR: File not found at {wb_path}")
        sys.exit(1)

    print(f"Loading workbook: {wb_path}...")
    wb = openpyxl.load_workbook(wb_path, read_only=True, data_only=True)
    print(f"Worksheets found: {wb.sheetnames}")

    if "Sheet16" not in wb.sheetnames:
        print("ERROR: Required worksheet 'Sheet16' not found in workbook.")
        sys.exit(1)

    print(f"\nInitiating Sheet16 Data Replacer (dry_run={args.dry_run})...")
    async with AsyncSessionLocal() as session:
        from app.core.auto_migrate import run_migrations_now
        await run_migrations_now(session)
        replacer = Sheet16DataReplacer(session, dry_run=args.dry_run)
        report = await replacer.execute(wb)

    wb.close()

    print("\n" + "=" * 70)
    print("           SHEET16 DATA REPLACEMENT & ARCHIVING REPORT")
    print("=" * 70)
    print(f"Batch ID:        {report['batch_id']}")
    print(f"Dry Run:         {report['dry_run']}")
    print(f"Status:          {report['status']}")
    print(f"Timestamp:       {report['timestamp']}")

    print("\n--- User Accounts & Deduplication ---")
    u_rep = report["users"]
    print(f"Canonical user accounts: {u_rep.get('canonical_users_count', 0)}")
    print(f"Duplicate accounts resolved: {u_rep.get('duplicates_resolved', 0)}")
    for d in u_rep.get("duplicate_details", []):
        print(f"  Merged: {d['email']} ({d['duplicate_id']} -> {d['canonical_id']})")

    print("\n--- Master Companies (Companies Worksheet) ---")
    c_rep = report["companies"]
    print(f"Master companies in sheet:   {c_rep.get('master_companies_in_sheet', 0)}")
    print(f"New companies created in DB: {c_rep.get('new_companies_created', 0)}")
    print(f"Total canonical companies:   {c_rep.get('total_canonical_companies', 0)}")

    print("\n--- Leads Worksheet Archiving ---")
    a_rep = report["leads_archive"]
    print(f"Rows archived into leads_archive: {a_rep.get('archived_rows', 0)}")
    print(f"Skipped empty rows:               {a_rep.get('skipped_empty_rows', 0)}")

    print("\n--- Active Contacts (Sheet16 Worksheet Only) ---")
    ct_rep = report["active_contacts"]
    print(f"Sheet16 source rows:       {ct_rep.get('sheet16_source_rows', 0)}")
    print(f"Valid contact rows:        {ct_rep.get('valid_contact_rows', 0)}")
    print(f"Invalid / skipped rows:    {ct_rep.get('invalid_rows', 0)}")
    print(f"Unique canonical contacts: {ct_rep.get('canonical_unique_contacts', 0)}")
    print(f"Duplicate rows in Sheet16: {ct_rep.get('sheet16_duplicate_rows', 0)}")
    print(f"Contacts newly inserted:   {ct_rep.get('contacts_inserted', 0)}")
    print(f"Contacts updated/relinked: {ct_rep.get('contacts_updated', 0)}")
    print(f"Old active contacts retired (archived): {ct_rep.get('retired_old_contacts', 0)}")

    print("\nPer-Salesperson allocation in Sheet16:")
    for sp, cnt in sorted(ct_rep.get("per_salesperson_counts", {}).items(), key=lambda x: x[1], reverse=True):
        print(f"  {sp}: {cnt} contacts")

    print("\n--- Historical Activities Extracted ---")
    d_rep = report["historical_demos"]
    f_rep = report["historical_followups"]
    print(f"Historical Demos created:      {d_rep.get('demos_created', 0)}")
    print(f"Historical Follow-ups created: {f_rep.get('followups_created', 0)}")

    print("\n--- Before vs After Active Contact Counts ---")
    pre = report.get("pre_migration_counts", {})
    post = report.get("post_migration_counts", {})
    print(f"{'Metric':<35} | {'Before':<10} | {'After':<10}")
    print("-" * 60)
    print(f"{'Total Active Contacts':<35} | {str(pre.get('total_active_contacts', 0)):<10} | {str(post.get('total_active_contacts', 0)):<10}")
    print(f"{'Total Distinct Companies':<35} | {str(pre.get('total_companies', 0)):<10} | {str(post.get('total_companies', 0)):<10}")
    print(f"{'Total Demos':<35} | {str(pre.get('total_demos', 0)):<10} | {str(post.get('total_demos', 0)):<10}")
    print(f"{'Total Follow-ups':<35} | {str(pre.get('total_followups', 0)):<10} | {str(post.get('total_followups', 0)):<10}")

    print("\nPer-user active contact breakdown:")
    pre_users = pre.get("per_user", {})
    post_users = post.get("per_user", {})
    all_emails = sorted(set(list(pre_users.keys()) + list(post_users.keys())))
    for em in all_emails:
        b_cnt = pre_users.get(em, {}).get("active_contacts", 0)
        a_cnt = post_users.get(em, {}).get("active_contacts", 0)
        name = post_users.get(em, {}).get("name") or pre_users.get(em, {}).get("name") or em
        print(f"  {name} ({em}): {b_cnt} -> {a_cnt}")

    print("=" * 70 + "\n")


if __name__ == "__main__":
    asyncio.run(main())
