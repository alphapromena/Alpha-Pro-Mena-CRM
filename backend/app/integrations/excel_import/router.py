"""
Excel Import Router — POST /api/v1/integrations/excel-import

Allows DATA_OPS and above to upload an .xlsx file and trigger the same
reconciliation pipeline used by the CLI import script. The uploading
user (e.g. Aseel) is NEVER assigned as the owner of imported contacts;
ownership comes only from the salesperson column in the workbook.

Endpoints
---------
POST /api/v1/integrations/excel-import
    Upload an .xlsx file for import.
    Fields:
        file (UploadFile): the workbook (.xlsx only, max 50 MB)
        sheet_name (str, optional): single sheet to process; default all
        dry_run (bool, optional): preview without writing; default false

GET /api/v1/integrations/excel-import/status
    Returns count of contacts/companies currently in the DB (for before/after comparison).
"""
from __future__ import annotations

import io
import structlog
from typing import Optional

from fastapi import APIRouter, Depends, File, Form, HTTPException, UploadFile, status
from sqlalchemy import text
from sqlalchemy.ext.asyncio import AsyncSession

from app.auth.dependencies import require_data_ops_or_above
from app.database import get_db
from app.models.user import User

logger = structlog.get_logger(__name__)

router = APIRouter(prefix="/integrations/excel-import", tags=["Excel Import"])

_MAX_FILE_BYTES = 50 * 1024 * 1024  # 50 MB


@router.post("", summary="Import contacts from an Excel workbook")
async def import_excel(
    file: UploadFile = File(..., description="Excel workbook (.xlsx)"),
    sheet_name: Optional[str] = Form(None, description="Single sheet name to import"),
    dry_run: bool = Form(False, description="Preview only — no database writes"),
    current_user: User = Depends(require_data_ops_or_above),
    db: AsyncSession = Depends(get_db),
):
    """Upload an .xlsx workbook and run the Gulf Leads reconciliation pipeline.

    - The importing user's identity is used only for audit logging.
    - Contact ownership comes from the salesperson column in the workbook.
    - The same idempotency and deduplication rules apply as the CLI import.
    """
    try:
        import openpyxl  # noqa: F401
    except ImportError:
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail="openpyxl is not installed on this server.",
        )

    # ── Validate file ─────────────────────────────────────────────────────
    if not file.filename or not file.filename.lower().endswith(".xlsx"):
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail="Only .xlsx files are supported.",
        )

    content = await file.read()
    if len(content) > _MAX_FILE_BYTES:
        raise HTTPException(
            status_code=status.HTTP_413_REQUEST_ENTITY_TOO_LARGE,
            detail=f"File exceeds the 50 MB limit ({len(content) // 1024 // 1024} MB received).",
        )

    # ── Parse workbook ────────────────────────────────────────────────────
    try:
        wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail=f"Could not open workbook: {exc}",
        )

    from app.imports.workbook_reader import (
        GULF_LEADS_SHEET_CONFIGS,
        SKIP_SHEETS,
        parse_sheet,
        read_companies_tab,
    )
    from app.imports.db_writer import run_import

    companies_set = read_companies_tab(wb)

    if sheet_name:
        if sheet_name not in wb.sheetnames:
            raise HTTPException(
                status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
                detail=f"Sheet '{sheet_name}' not found. Available: {wb.sheetnames}",
            )
        sheets_to_process = [sheet_name]
    else:
        sheets_to_process = [
            s for s in wb.sheetnames
            if s not in SKIP_SHEETS and s in GULF_LEADS_SHEET_CONFIGS
        ]

    all_rows = []
    sheets_processed = []
    for sname in sheets_to_process:
        config = GULF_LEADS_SHEET_CONFIGS.get(sname)
        if not config:
            continue
        ws = wb[sname]
        rows = parse_sheet(ws, config, companies_set)
        all_rows.extend(rows)
        sheets_processed.append({"sheet": sname, "rows": len(rows)})

    wb.close()

    if not all_rows:
        return {
            "message": "No processable rows found.",
            "sheets_processed": sheets_processed,
            "report": None,
        }

    # ── Run import ────────────────────────────────────────────────────────
    report = await run_import(db, all_rows, dry_run=dry_run)

    logger.info(
        "excel_import.completed",
        user_id=str(current_user.id),
        file_name=file.filename,
        dry_run=dry_run,
        inserted=report.total_inserted,
        updated=report.total_updated,
        skipped=report.total_skipped,
    )

    return {
        "message": "Dry-run completed — no records written." if dry_run else "Import completed.",
        "dry_run": dry_run,
        "uploaded_by": str(current_user.id),
        "sheets_processed": sheets_processed,
        "report": report.to_dict(),
    }


@router.get("/status", summary="Current contact/company counts")
async def import_status(
    current_user: User = Depends(require_data_ops_or_above),
    db: AsyncSession = Depends(get_db),
):
    """Return the current count of contacts and companies for before/after comparison."""
    contacts = (await db.execute(
        text("SELECT COUNT(*) FROM contacts WHERE deleted_at IS NULL")
    )).scalar()
    companies = (await db.execute(
        text("SELECT COUNT(*) FROM companies WHERE deleted_at IS NULL")
    )).scalar()
    return {"contacts": contacts, "companies": companies}
